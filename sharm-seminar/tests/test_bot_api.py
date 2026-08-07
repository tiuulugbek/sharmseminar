import datetime
import hashlib
import hmac
import importlib.util
import json
import os
import sqlite3
import tempfile
import time
import unittest
from urllib.parse import urlencode


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SPEC = importlib.util.spec_from_file_location("seminar_server", os.path.join(ROOT, "server.py"))
server = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(server)


class BotApiTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        server.DB_PATH = os.path.join(self.tmp.name, "test.db")
        server.SEED_P = os.path.join(self.tmp.name, "missing.json")
        server._db_initialized = True
        os.environ["BOT_API_TOKEN"] = "test-token"
        os.environ["SECRET"] = "test-secret"
        os.environ["ADMIN_IDS"] = "555"
        os.environ["BOT_TOKEN"] = "123:TESTBOTTOKEN"
        os.environ["PANEL_ADMIN_IDS"] = ""
        server.init_db()
        self.client = server.app.test_client()
        self.headers = {"X-Bot-Token": "test-token"}

    def panel(self, pid):
        """Maxfiy kod (pasport) bilan kirgan panel klienti."""
        row = self.one("SELECT passport_series,passport_number FROM participants WHERE id=?", pid)
        client = server.app.test_client()
        response = client.post("/api/auth/login", json={"code": (row[0] or "") + (row[1] or "")})
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        return client

    def sql(self, statement, *params):
        con = sqlite3.connect(server.DB_PATH)
        con.execute(statement, params)
        con.commit(); con.close()

    def _all(self, statement, *params):
        con = sqlite3.connect(server.DB_PATH)
        rows = con.execute(statement, params).fetchall()
        con.close()
        return rows

    def one(self, statement, *params):
        con = sqlite3.connect(server.DB_PATH)
        row = con.execute(statement, params).fetchone()
        con.close()
        return row

    def tearDown(self):
        self.tmp.cleanup()

    def post(self, path, data):
        return self.client.post(path, json=data, headers=self.headers)

    def register(self, fio, number, telegram_id=""):
        response = self.post("/api/bot/register", {
            "fio": fio, "passport_series": "FA", "passport_number": number,
            "dob": "01-02-2000", "xona_turi": "2 kishilik",
            "telegram_id": telegram_id,
        })
        self.assertEqual(response.status_code, 201, response.get_data(as_text=True))
        return response.get_json()["participant"]

    def test_requires_token(self):
        self.assertEqual(self.client.get("/api/bot/stats").status_code, 401)

    def test_legacy_schema_migration_is_idempotent(self):
        legacy = os.path.join(self.tmp.name, "legacy.db")
        con = sqlite3.connect(legacy)
        con.execute("CREATE TABLE participants(id TEXT PRIMARY KEY,fio TEXT,jinsi TEXT,fuqarolik TEXT,"
                    "xona_turi TEXT,xona_guruhi TEXT,kelish TEXT,grp INTEGER,leader INTEGER DEFAULT 0,"
                    "room TEXT DEFAULT '',branch TEXT DEFAULT '',telegram TEXT DEFAULT '',roles TEXT DEFAULT '[]')")
        con.execute("CREATE TABLE checkins(pid TEXT,checkpoint TEXT,ts TEXT,PRIMARY KEY(pid,checkpoint))")
        con.execute("CREATE TABLE settings(key TEXT PRIMARY KEY,value TEXT)")
        con.execute("INSERT INTO participants(id,fio) VALUES('OLD-1','Old Person')")
        con.commit(); con.close()
        server.DB_PATH = legacy
        server.init_db(); server.init_db()
        con = sqlite3.connect(legacy)
        cols = {r[1] for r in con.execute("PRAGMA table_info(participants)")}
        self.assertIn("passport_series", cols)
        self.assertEqual(con.execute("SELECT fio FROM participants WHERE id='OLD-1'").fetchone()[0], "Old Person")
        con.close()

    def test_register_find_leading_zero_and_duplicate(self):
        person = self.register("Karimov Ali", "0012345")
        found = self.post("/api/bot/find", {"series": "FA12345"}).get_json()["participants"]
        self.assertEqual(found[0]["id"], person["id"])
        duplicate = self.post("/api/bot/register", {
            "fio": "Duplicate", "passport_series": "FA", "passport_number": "12345"})
        self.assertEqual(duplicate.status_code, 409)

    def test_dob_formats_and_telegram_conflict(self):
        one = self.register("One Person", "0000001")
        two = self.register("Two Person", "0000002")
        found = self.post("/api/bot/find", {"dob": "01.02.2000"}).get_json()["participants"]
        self.assertEqual(len(found), 2)
        self.assertEqual(self.post("/api/bot/link", {
            "id": one["id"], "telegram_id": "777"}).status_code, 200)
        self.assertEqual(self.post("/api/bot/link", {
            "id": two["id"], "telegram_id": "777"}).status_code, 409)

    def test_roommate_recipients_and_stats(self):
        one = self.register("One Person", "0000001", "101")
        self.register("Two Person", "0000002", "102")
        linked = self.post("/api/bot/roommate", {
            "series": "FA1", "roommate_series": "FA2"}).get_json()
        self.assertEqual(linked["status"], "linked")
        recipients = self.client.get("/api/bot/recipients", headers=self.headers).get_json()["recipients"]
        self.assertEqual(len(recipients), 2)
        self.assertEqual(recipients[0]["roommates"], ["Two Person"])
        stats = self.client.get("/api/bot/stats", headers=self.headers).get_json()
        self.assertEqual(stats["total"], 2)
        self.assertEqual(stats["rooms_ready"], 1)

    # ------------------------------------------------------ tokens & QR pages
    def test_token_is_passport_derived_stable_and_hides_the_passport(self):
        person = self.register("Token Person", "0000123")
        token = person["token"]
        self.assertEqual(len(token), 16)
        self.assertEqual(token, server.make_token("FA", "0000123", person["id"]))
        # Re-running the migration must not hand out a different token.
        server.init_db()
        self.assertEqual(self.one("SELECT token FROM participants WHERE id=?", person["id"])[0], token)
        # The public page opens by token and by id, and leaks no passport.
        by_token = self.client.get(f"/api/p/{token}")
        by_id = self.client.get(f"/api/p/{person['id']}")
        self.assertEqual(by_token.status_code, 200)
        self.assertEqual(by_id.status_code, 200)
        self.assertEqual(by_token.get_json()["participant"]["id"], person["id"])
        self.assertNotIn("0000123", by_token.get_data(as_text=True))
        self.assertEqual(self.client.get("/api/p/nosuchtoken").status_code, 404)

    def test_token_differs_per_secret(self):
        person = self.register("Secret Person", "0000124")
        os.environ["SECRET"] = "another-secret"
        self.assertNotEqual(server.make_token("FA", "0000124", person["id"]), person["token"])
        os.environ["SECRET"] = "test-secret"

    # ------------------------------------------------------------------ roles
    def build_group(self):
        """Group 1 with a leader and two members, group 2 with one member."""
        leader = self.register("Leader One", "0000001", "9001")
        member = self.register("Member One", "0000002", "9100")
        other = self.register("Other Group", "0000003", "9200")
        self.sql("UPDATE participants SET grp=1,leader=1 WHERE id=?", leader["id"])
        self.sql("UPDATE participants SET grp=1 WHERE id=?", member["id"])
        self.sql("UPDATE participants SET grp=2 WHERE id=?", other["id"])
        return leader, member, other

    def whoami(self, telegram_id):
        return self.client.get(f"/api/bot/whoami?telegram_id={telegram_id}",
                               headers=self.headers).get_json()

    def test_panel_role_carries_into_the_bot(self):
        """Rollar bitta joyda — bazada — boshqariladi, .env bilan bo'linmaydi."""
        leader, member, other = self.build_group()
        self.assertEqual(self.whoami("9001")["role"], "leader")

        # Rahbar botda ham barcha guruh bilan ishlaydi.
        self.sql("UPDATE participants SET panel_role='manager' WHERE id=?", leader["id"])
        self.assertEqual(self.whoami("9001")["role"], "admin")

        self.sql("UPDATE participants SET panel_role='admin' WHERE id=?", member["id"])
        self.assertEqual(self.whoami("9100")["role"], "admin")

        # Pastga tushirilsa ham darhol kuchga kiradi.
        self.sql("UPDATE participants SET panel_role=NULL WHERE id=?", member["id"])
        self.assertEqual(self.whoami("9100")["role"], "member")

    def test_whoami_ranks_admin_over_leader_over_member(self):
        leader, member, _ = self.build_group()
        self.assertEqual(self.whoami("555")["role"], "admin")
        self.assertEqual(self.whoami("9001")["role"], "leader")
        self.assertEqual(self.whoami("9100")["role"], "member")
        self.assertEqual(self.whoami("404040")["role"], "guest")
        self.assertEqual(self.whoami("9001")["group_name"], "Sazanchik")
        # An admin who is also on the participant list stays an admin.
        self.sql("UPDATE participants SET telegram_id='555' WHERE id=?", member["id"])
        self.assertEqual(self.whoami("555")["role"], "admin")

    # --------------------------------------------------------------- check-in
    def test_checkin_is_limited_to_admins_and_the_leaders_own_group(self):
        leader, member, other = self.build_group()
        ok = self.post("/api/bot/checkin", {
            "token": member["token"], "checkpoint": "seminar", "by_telegram_id": "9001"})
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(ok.get_json()["participant"]["id"], member["id"])

        for who, payload in (("leader over another group", {"token": other["token"], "by_telegram_id": "9001"}),
                             ("member", {"token": member["token"], "by_telegram_id": "9100"}),
                             ("stranger", {"token": member["token"], "by_telegram_id": "404040"})):
            response = self.post("/api/bot/checkin", dict(payload, checkpoint="seminar"))
            self.assertEqual(response.status_code, 403, who)

        self.assertEqual(self.post("/api/bot/checkin", {
            "token": other["token"], "checkpoint": "seminar", "by_telegram_id": "555"}).status_code, 200)
        self.assertEqual(self.post("/api/bot/checkin", {
            "token": member["token"], "checkpoint": "nosuch", "by_telegram_id": "555"}).status_code, 400)

    def test_leader_scope_all_lets_leaders_scan_anybody(self):
        _, _, other = self.build_group()
        con = sqlite3.connect(server.DB_PATH)
        con.execute("UPDATE settings SET value='\"all\"' WHERE key='leader_scope'")
        con.commit(); con.close()
        self.assertEqual(self.post("/api/bot/checkin", {
            "token": other["token"], "checkpoint": "seminar", "by_telegram_id": "9001"}).status_code, 200)

    def set_checkpoints(self, checkpoints):
        self.sql("UPDATE settings SET value=? WHERE key='checkpoints'",
                 json.dumps(checkpoints, ensure_ascii=False))

    def scan(self, token, checkpoint, by="555", **extra):
        return self.post("/api/bot/checkin",
                         dict({"token": token, "checkpoint": checkpoint,
                               "by_telegram_id": by}, **extra))

    def test_the_same_person_is_only_recorded_once_per_checkpoint(self):
        _, member, _ = self.build_group()
        first = self.scan(member["token"], "seminar").get_json()
        self.assertEqual(first["status"], "ok")

        again = self.scan(member["token"], "seminar").get_json()
        self.assertEqual(again["status"], "already")
        self.assertEqual(again["ts"], first["ts"])   # vaqt qayta yozilmaydi
        self.assertEqual(self.one("SELECT COUNT(*) FROM checkins WHERE pid=?", member["id"])[0], 1)

        # Admin ataylab qayta yozmoqchi bo'lsa — force bilan mumkin.
        forced = self.scan(member["token"], "seminar", force=True).get_json()
        self.assertEqual(forced["status"], "ok")

    def test_a_checkpoint_only_accepts_scans_inside_its_time_window(self):
        _, member, _ = self.build_group()
        now = server._event_now().replace(tzinfo=None)
        fmt = "%Y-%m-%dT%H:%M"
        self.set_checkpoints([
            {"key": "aeroport", "label": "Aeroport"},
            {"key": "past", "label": "Mehmonxona",
             "ends_at": (now - datetime.timedelta(days=1)).strftime(fmt)},
            {"key": "future", "label": "Yahta",
             "starts_at": (now + datetime.timedelta(days=1)).strftime(fmt)},
        ])
        self.assertEqual(self.scan(member["token"], "past").get_json()["status"], "closed")
        self.assertEqual(self.scan(member["token"], "future").get_json()["status"], "not_open")
        self.assertEqual(self.one("SELECT COUNT(*) FROM checkins WHERE pid=?", member["id"])[0], 0)

        # Vaqti kelmagan nuqtani admin force bilan ocha oladi.
        self.assertEqual(self.scan(member["token"], "future", force=True).get_json()["status"], "ok")

    def test_skipping_an_earlier_checkpoint_warns_but_never_blocks(self):
        _, member, _ = self.build_group()
        self.set_checkpoints([{"key": "aeroport", "label": "Aeroport"},
                              {"key": "mehmonxona", "label": "Mehmonxona"},
                              {"key": "yahta", "label": "Yahta"}])
        # Aeroport va mehmonxonasiz to'g'ridan-to'g'ri yahtada belgilanadi.
        body = self.scan(member["token"], "yahta").get_json()
        self.assertEqual(body["status"], "ok")
        self.assertEqual([w["key"] for w in body["warnings"]], ["aeroport", "mehmonxona"])
        # O'tkazib yuborilgani o'tkazib yuborilganicha qoladi.
        self.assertEqual(sorted(body["checkins"]), ["yahta"])

        later = self.scan(member["token"], "mehmonxona").get_json()
        self.assertEqual(later["status"], "ok")
        self.assertEqual([w["key"] for w in later["warnings"]], ["aeroport"])

    def test_checkin_timestamps_use_event_time_not_server_time(self):
        _, member, _ = self.build_group()
        self.sql("INSERT INTO settings(key,value) VALUES('tz_offset','9') "
                 "ON CONFLICT(key) DO UPDATE SET value='9'")
        stamped = self.scan(member["token"], "seminar").get_json()["ts"]
        expected = (datetime.datetime.now(datetime.timezone.utc)
                    + datetime.timedelta(hours=9)).strftime("%H:%M")
        self.assertEqual(stamped, expected)

    def test_unchecking_clears_the_mark_so_it_can_be_scanned_again(self):
        _, member, _ = self.build_group()
        self.scan(member["token"], "seminar")
        self.assertEqual(self.scan(member["token"], "seminar",
                                   on=False).get_json()["status"], "removed")
        self.assertEqual(self.scan(member["token"], "seminar").get_json()["status"], "ok")

    # -------------------------------------------------------------- messaging
    def send(self, telegram_id, scope, text, value=None):
        return self.post("/api/bot/message", {
            "from_telegram_id": telegram_id, "scope": scope,
            "value": value, "text": text})

    def test_admin_broadcast_reaches_everyone_and_replies_come_back_to_admin(self):
        leader, member, other = self.build_group()
        sent = self.send("555", "all", "hammaga")
        self.assertEqual(sent.status_code, 200)
        self.assertEqual({r["id"] for r in sent.get_json()["recipients"]},
                         {leader["id"], member["id"], other["id"]})

        msg_id = sent.get_json()["message_id"]
        answer = self.post("/api/bot/reply", {
            "from_telegram_id": "9100", "parent_msg_id": msg_id, "text": "javob"})
        self.assertEqual(answer.status_code, 200)
        body = answer.get_json()
        self.assertEqual(body["to_telegram_id"], "555")
        self.assertIn(member["id"], body["sender_label"])
        self.assertIn("Member One", body["sender_label"])

    def test_leader_may_only_write_to_their_own_group(self):
        leader, member, other = self.build_group()
        mine = self.send("9001", "group", "guruhimga")
        self.assertEqual(mine.status_code, 200)
        self.assertEqual([r["id"] for r in mine.get_json()["recipients"]], [member["id"]])

        self.assertEqual(self.send("9001", "group", "x", value=2).status_code, 403)
        self.assertEqual(self.send("9001", "all", "x").status_code, 403)
        self.assertEqual(self.send("9001", "one", "x", value=other["id"]).status_code, 403)

        answer = self.post("/api/bot/reply", {
            "from_telegram_id": "9100", "parent_msg_id": mine.get_json()["message_id"],
            "text": "javob"})
        self.assertEqual(answer.get_json()["to_telegram_id"], "9001")

    def test_member_can_only_ask_their_leader_and_gets_the_answer_back(self):
        leader, member, other = self.build_group()
        self.assertEqual(self.send("9100", "all", "x").status_code, 403)
        self.assertEqual(self.send("9100", "group", "x").status_code, 403)
        self.assertEqual(self.send("9100", "one", "x", value=other["id"]).status_code, 403)

        question = self.send("9100", "leader", "savol")
        self.assertEqual(question.status_code, 200)
        self.assertEqual([r["telegram_id"] for r in question.get_json()["recipients"]], ["9001"])

        answer = self.post("/api/bot/reply", {
            "from_telegram_id": "9001", "parent_msg_id": question.get_json()["message_id"],
            "text": "javob"})
        self.assertEqual(answer.get_json()["to_telegram_id"], "9100")

    def test_outsiders_cannot_reply_and_delivery_is_recorded(self):
        leader, member, other = self.build_group()
        mine = self.send("9001", "group", "guruhimga").get_json()
        self.assertEqual(self.post("/api/bot/reply", {
            "from_telegram_id": "9200", "parent_msg_id": mine["message_id"],
            "text": "x"}).status_code, 403)
        self.assertEqual(self.post("/api/bot/reply", {
            "from_telegram_id": "9100", "parent_msg_id": 99999, "text": "x"}).status_code, 404)

        self.post("/api/bot/message/sent", {"message_id": mine["message_id"], "results": [
            {"id": member["id"], "telegram_id": "9100", "telegram_msg_id": "42", "ok": True}]})
        self.assertEqual(self.one("SELECT telegram_msg_id,status FROM msg_targets WHERE msg_id=?",
                                  mine["message_id"]), ("42", "sent"))
        self.assertEqual(self.one("SELECT delivered FROM messages WHERE id=?", mine["message_id"])[0], 1)

    def test_member_questions_are_copied_to_admins_when_enabled(self):
        self.build_group()
        con = sqlite3.connect(server.DB_PATH)
        con.execute("UPDATE settings SET value='true' WHERE key='copy_member_questions'")
        con.commit(); con.close()
        self.assertEqual(self.send("9100", "leader", "savol").get_json()["copy_to"], ["555"])

    # ------------------------------------------------------- WebApp initData
    def init_data(self, telegram_id, bot_token="123:TESTBOTTOKEN", auth_date=None):
        fields = {"auth_date": str(auth_date or int(time.time())),
                  "user": json.dumps({"id": telegram_id, "username": "tester"})}
        check = "\n".join(f"{k}={fields[k]}" for k in sorted(fields))
        secret = hmac.new(b"WebAppData", bot_token.encode(), hashlib.sha256).digest()
        fields["hash"] = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
        return urlencode(fields)

    def test_webapp_rejects_forged_or_stale_init_data(self):
        self.build_group()
        self.assertEqual(self.client.post("/api/webapp/whoami", json={}).status_code, 401)
        self.assertEqual(self.client.post("/api/webapp/whoami", json={
            "init_data": "user=%7B%22id%22%3A555%7D&hash=deadbeef"}).status_code, 401)
        self.assertEqual(self.client.post("/api/webapp/whoami", json={
            "init_data": self.init_data(9001, bot_token="999:WRONG")}).status_code, 401)
        self.assertEqual(self.client.post("/api/webapp/whoami", json={
            "init_data": self.init_data(9001, auth_date=int(time.time()) - 90000)}).status_code, 401)

    def test_webapp_checkin_trusts_only_the_signed_identity(self):
        leader, member, other = self.build_group()
        signed_leader = self.init_data(9001)
        response = self.client.post("/api/webapp/whoami", json={"init_data": signed_leader})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["role"], "leader")

        ok = self.client.post("/api/webapp/checkin", json={
            "init_data": signed_leader, "token": member["token"], "checkpoint": "seminar"})
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(ok.get_json()["participant"]["id"], member["id"])

        # A signed member cannot check anybody in, not even themselves.
        denied = self.client.post("/api/webapp/checkin", json={
            "init_data": self.init_data(9100), "token": member["token"], "checkpoint": "seminar"})
        self.assertEqual(denied.status_code, 403)
        # Nor can a leader reach another group.
        self.assertEqual(self.client.post("/api/webapp/checkin", json={
            "init_data": signed_leader, "token": other["token"],
            "checkpoint": "seminar"}).status_code, 403)

    # --------------------------------------------------------------- rosters
    def test_group_endpoints_report_leaders_and_attendance(self):
        leader, member, _ = self.build_group()
        self.post("/api/bot/checkin", {"token": member["token"], "checkpoint": "seminar",
                                       "by_telegram_id": "555"})
        groups = self.client.get("/api/bot/groups", headers=self.headers).get_json()["groups"]
        first = next(g for g in groups if g["id"] == 1)
        self.assertEqual(first["name"], "Sazanchik")
        self.assertEqual(first["leader"]["id"], leader["id"])
        self.assertEqual((first["total"], first["checked_in"]), (2, 1))

        roster = self.client.get("/api/bot/group/1", headers=self.headers).get_json()
        self.assertEqual(roster["members"][0]["id"], leader["id"])  # leader first
        self.assertIn("seminar", next(m for m in roster["members"]
                                      if m["id"] == member["id"])["checkins"])

    def test_saving_groups_moves_the_leader_flag(self):
        leader, member, _ = self.build_group()
        os.environ["PANEL_ADMIN_IDS"] = leader["id"]
        self.panel(leader["id"]).post("/api/groups", json={"groups": [
            {"id": 1, "name": "Sazanchik", "leader": member["id"]}]})
        self.assertEqual(self.one("SELECT leader FROM participants WHERE id=?", member["id"])[0], 1)
        self.assertEqual(self.one("SELECT leader FROM participants WHERE id=?", leader["id"])[0], 0)

    # --------------------------------------------------- panelga kirish (kod)
    def test_login_accepts_the_passport_in_several_shapes(self):
        person = self.register("Login Person", "0012345")
        for code in ("FA0012345", "fa0012345", "FA 0012345", "0012345", "12345", "FA12345"):
            client = server.app.test_client()
            response = client.post("/api/auth/login", json={"code": code})
            self.assertEqual(response.status_code, 200, code)
            self.assertEqual(response.get_json()["user"]["id"], person["id"], code)

    def test_login_rejects_wrong_short_and_ambiguous_codes(self):
        self.register("One", "0000001")
        self.register("Two", "0000002")
        self.assertEqual(self.client.post("/api/auth/login", json={"code": "FA9999999"}).status_code, 401)
        self.assertEqual(self.client.post("/api/auth/login", json={"code": "12"}).status_code, 401)
        self.assertEqual(self.client.post("/api/auth/login", json={}).status_code, 401)

    def test_session_cookie_cannot_be_forged_and_logout_clears_it(self):
        person = self.register("Cookie Person", "0000321")
        client = self.panel(person["id"])
        self.assertEqual(client.get("/api/auth/me").status_code, 200)
        client.post("/api/auth/logout")
        self.assertEqual(client.get("/api/auth/me").status_code, 401)

        forged = server.app.test_client()
        forged.set_cookie(server.SESSION_COOKIE, f"{person['id']}|admin|1|deadbeef")
        self.assertEqual(forged.get("/api/auth/me").status_code, 401)
        # A valid signature for "member" must not be replayable as "admin".
        signed = server._sign_session(person["id"], "member")
        tampered = server.app.test_client()
        tampered.set_cookie(server.SESSION_COOKIE, signed.replace("|member|", "|admin|"))
        self.assertEqual(tampered.get("/api/auth/me").status_code, 401)

    def test_role_comes_from_the_database_not_from_the_cookie(self):
        person = self.register("Climber", "0000444")
        client = self.panel(person["id"])
        self.assertEqual(client.get("/api/auth/me").get_json()["user"]["role"], "member")
        self.sql("UPDATE participants SET panel_role='manager' WHERE id=?", person["id"])
        self.assertEqual(client.get("/api/auth/me").get_json()["user"]["role"], "manager")

    def test_bot_admin_telegram_id_also_opens_the_panel_as_admin(self):
        person = self.register("Bot Admin", "0000555", "555")  # 555 is in ADMIN_IDS
        self.assertEqual(self.panel(person["id"]).get("/api/auth/me")
                         .get_json()["user"]["role"], "admin")

    def test_bootstrap_shows_only_what_the_role_may_see(self):
        leader, member, other = self.build_group()
        admin = self.register("Panel Admin", "0000777")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        self.sql("UPDATE participants SET panel_role='manager' WHERE id=?", other["id"])

        self.assertEqual(self.client.get("/api/bootstrap").status_code, 401)

        seen = {}
        for pid in (admin["id"], other["id"], leader["id"], member["id"]):
            data = self.panel(pid).get("/api/bootstrap").get_json()
            seen[data["user"]["role"]] = data

        self.assertEqual(len(seen["admin"]["participants"]), 4)
        self.assertEqual(len(seen["manager"]["participants"]), 4)
        self.assertEqual({p["id"] for p in seen["leader"]["participants"]},
                         {leader["id"], member["id"]})
        self.assertEqual([p["id"] for p in seen["member"]["participants"]], [member["id"]])

        # Passports stay with managers and above; a leader sees none but their own.
        leaked = [p for p in seen["leader"]["participants"]
                  if p["id"] != leader["id"] and p.get("passport_number")]
        self.assertEqual(leaked, [])
        self.assertTrue(any(p.get("passport_number") for p in seen["manager"]["participants"]))

    def test_writes_are_limited_by_role(self):
        leader, member, other = self.build_group()
        admin = self.register("Panel Admin", "0000778")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        self.sql("UPDATE participants SET panel_role='manager' WHERE id=?", other["id"])

        as_admin, as_manager = self.panel(admin["id"]), self.panel(other["id"])
        as_leader, as_member = self.panel(leader["id"]), self.panel(member["id"])

        # Settings: admin only.
        self.assertEqual(as_admin.post("/api/meta", json={"title": "x"}).status_code, 200)
        for client in (as_manager, as_leader, as_member):
            self.assertEqual(client.post("/api/meta", json={"title": "x"}).status_code, 403)

        # Xona raqami va til: guruh mas'ulidan boshlab, faqat o'z guruhida.
        patch = {"id": member["id"], "patch": {"room": "101"}}
        self.assertEqual(as_manager.post("/api/participant", json=patch).status_code, 200)
        self.assertEqual(as_leader.post("/api/participant", json=patch).status_code, 200)
        self.assertEqual(as_member.post("/api/participant", json=patch).status_code, 403)
        # Mas'ul begona guruhga va ruxsat etilmagan maydonga tegolmaydi.
        self.assertEqual(as_leader.post("/api/participant", json={
            "id": other["id"], "patch": {"room": "9"}}).status_code, 403)
        self.assertEqual(as_leader.post("/api/participant", json={
            "id": member["id"], "patch": {"xona_guruhi": "D40"}}).status_code, 403)
        self.assertEqual(as_leader.post("/api/participant", json={
            "id": member["id"], "patch": {"staff": True}}).status_code, 403)

        # Check-in: a leader only inside their own group, a member never.
        own = {"id": member["id"], "checkpoint": "seminar", "on": True}
        outside = {"id": other["id"], "checkpoint": "seminar", "on": True}
        self.assertEqual(as_leader.post("/api/checkin", json=own).status_code, 200)
        self.assertEqual(as_leader.post("/api/checkin", json=outside).status_code, 403)
        self.assertEqual(as_manager.post("/api/checkin", json=outside).status_code, 200)
        self.assertEqual(as_member.post("/api/checkin", json=own).status_code, 403)

    def test_only_an_admin_may_change_groups(self):
        """Taqsimot bir marta to'g'rilangach, tasodifan aralashib ketmasin."""
        leader, member, other = self.build_group()
        admin = self.register("Panel Admin", "0000781")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        self.sql("UPDATE participants SET panel_role='manager' WHERE id=?", other["id"])
        as_admin, as_manager = self.panel(admin["id"]), self.panel(other["id"])

        self.assertEqual(as_admin.post("/api/participant", json={
            "id": member["id"], "patch": {"group": 3}}).status_code, 200)
        for field in ({"group": 4}, {"leader": True}):
            response = as_manager.post("/api/participant", json={"id": member["id"], "patch": field})
            self.assertEqual(response.status_code, 403, field)
        self.assertEqual(self.one("SELECT grp FROM participants WHERE id=?", member["id"])[0], 3)

        # Xona va til rahbarga ochiq qoladi — cheklov faqat guruhga.
        self.assertEqual(as_manager.post("/api/participant", json={
            "id": member["id"], "patch": {"room": "214", "lang": "en"}}).status_code, 200)

        # Ommaviy qayta taqsimlash ham faqat admin qo'lida.
        items = {"items": [{"id": member["id"], "group": 5}]}
        self.assertEqual(as_manager.post("/api/participants/bulk", json=items).status_code, 403)
        self.assertEqual(as_admin.post("/api/participants/bulk", json=items).status_code, 200)
        self.assertEqual(self.one("SELECT grp FROM participants WHERE id=?", member["id"])[0], 5)

    def test_moving_someone_between_rooms_updates_both_sides(self):
        leader, member, other = self.build_group()
        admin = self.register("Panel Admin", "0000780")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        panel = self.panel(admin["id"])
        self.sql("UPDATE participants SET xona_guruhi='D07' WHERE id IN (?,?)",
                 leader["id"], member["id"])
        self.sql("UPDATE participants SET xona_guruhi='D08', roommate_series='FA9' WHERE id=?",
                 other["id"])

        # Boshqa blokka ko'chirish: eski xonadosh yolg'iz qoladi, yangisi qo'shiladi.
        self.assertEqual(panel.post("/api/participant", json={
            "id": member["id"], "patch": {"xona_guruhi": "d08"}}).status_code, 200)
        page = self.client.get(f"/api/p/{member['token']}").get_json()
        self.assertEqual(page["participant"]["xona_guruhi"], "D08")   # bosh harfga keltiriladi
        self.assertEqual(page["roommates"], ["Other Group"])
        self.assertEqual(self.client.get(f"/api/p/{leader['token']}").get_json()["roommates"], [])

        # Eski pasport-asosidagi sheriklik tozalanadi, u endi to'g'ri kelmaydi.
        self.assertIsNone(self.one("SELECT roommate_series FROM participants WHERE id=?",
                                   member["id"])[0])
        # Guruh mas'uli xona ko'chira olmaydi.
        self.assertEqual(self.panel(leader["id"]).post("/api/participant", json={
            "id": member["id"], "patch": {"xona_guruhi": "D07"}}).status_code, 403)

    def test_replacing_a_participant_keeps_the_seat_and_clears_the_person(self):
        leader, member, other = self.build_group()
        admin = self.register("Panel Admin", "0000796")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        server.DOCS_FILES = os.path.join(server.DOCS_DIR, "_files")
        panel = self.panel(admin["id"])
        self.sql("UPDATE participants SET xona_guruhi='D07',room='214' WHERE id=?", member["id"])
        self.upload(panel, [f"{member['id']} voucher.pdf"])
        self.post("/api/bot/checkin", {"token": member["token"], "checkpoint": "seminar",
                                       "by_telegram_id": "555"})

        response = panel.post("/api/participant/replace", json={
            "id": member["id"], "fio": "Yangi Odam", "passport": "FA9990002",
            "dob": "15.03.1995", "jinsi": "MR", "fuqarolik": "UZBEKISTAN"})
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        body = response.get_json()
        self.assertEqual(body["was"]["fio"], "Member One")
        self.assertEqual(body["documents_removed"], 1)

        row = self.one("SELECT fio,grp,xona_guruhi,room,telegram_id,leader,token,"
                       "passport_number FROM participants WHERE id=?", member["id"])
        fio, grp, block, room, telegram, is_leader, token, passport = row
        self.assertEqual(fio, "Yangi Odam")
        self.assertEqual((grp, block, room), (1, "D07", "214"))      # o'rin saqlanadi
        self.assertIsNone(telegram)                                   # odam almashdi
        self.assertEqual(is_leader, 0)
        self.assertEqual(passport, "9990002")
        self.assertNotEqual(token, member["token"])                   # QR yangi
        self.assertEqual(self.one("SELECT COUNT(*) FROM documents WHERE pid=?",
                                  member["id"])[0], 0)
        self.assertEqual(self.one("SELECT COUNT(*) FROM checkins WHERE pid=?",
                                  member["id"])[0], 0)
        self.assertEqual(len(self._all("SELECT id FROM participants")), 4)   # soni o'zgarmaydi

    def test_replacement_refuses_a_passport_somebody_else_already_has(self):
        leader, member, _ = self.build_group()
        admin = self.register("Panel Admin", "0000797")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        panel = self.panel(admin["id"])
        # Seriyasi alohida saqlangan pasport ham topilishi kerak.
        self.sql("UPDATE participants SET passport_series='77',passport_number='3408359' "
                 "WHERE id=?", leader["id"])

        for code in ("773408359", "3408359", "77 3408359"):
            taken = panel.post("/api/participant/replace", json={
                "id": member["id"], "fio": "Test Kimdir", "passport": code})
            self.assertEqual(taken.status_code, 409, code)
            self.assertEqual(taken.get_json()["error"], "passport_already_exists")

        self.assertEqual(panel.post("/api/participant/replace", json={
            "id": member["id"], "fio": "", "passport": "FA1111111"}).status_code, 400)
        # Hech narsa o'zgarmagan.
        self.assertEqual(self.one("SELECT fio FROM participants WHERE id=?",
                                  member["id"])[0], "Member One")

        # Faqat admin qila oladi.
        self.assertEqual(self.panel(leader["id"]).post("/api/participant/replace", json={
            "id": member["id"], "fio": "X Y", "passport": "FA1111111"}).status_code, 403)

    def test_unlinking_frees_a_slot_for_the_real_person(self):
        leader, member, other = self.build_group()
        admin = self.register("Panel Admin", "0000779")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        panel = self.panel(admin["id"])

        # Noto'g'ri odam ro'yxatdan o'tkazib qo'ygan: 9100 aslida boshqasi edi.
        self.assertEqual(self.whoami("9100")["id"], member["id"])
        response = panel.post("/api/participant/unlink", json={"id": member["id"]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["was"]["telegram_id"], "9100")
        self.assertEqual(self.whoami("9100")["role"], "guest")

        # Guruh, xona va QR token joyida qoladi — beyjik ishlashda davom etadi.
        row = self.one("SELECT grp,token FROM participants WHERE id=?", member["id"])
        self.assertEqual(row, (1, member["token"]))

        # Endi haqiqiy egasi o'zini bog'lay oladi.
        self.assertEqual(self.post("/api/bot/link", {
            "id": member["id"], "telegram_id": "9300"}).status_code, 200)
        self.assertEqual(self.whoami("9300")["id"], member["id"])

        # Guruh mas'uli o'z guruhida uza oladi, begonaga tegolmaydi.
        self.post("/api/bot/link", {"id": member["id"], "telegram_id": "9400"})
        self.assertEqual(self.panel(leader["id"]).post(
            "/api/participant/unlink", json={"id": member["id"]}).status_code, 200)
        self.assertEqual(self.panel(leader["id"]).post(
            "/api/participant/unlink", json={"id": other["id"]}).status_code, 403)
        # Oddiy ishtirokchiga va imzosizga yopiq.
        self.assertEqual(self.panel(other["id"]).post(
            "/api/participant/unlink", json={"id": member["id"]}).status_code, 403)
        self.assertEqual(self.client.post(
            "/api/participant/unlink", json={"id": member["id"]}).status_code, 401)

    def test_scanner_roster_shows_who_is_still_missing(self):
        leader, member, other = self.build_group()
        self.post("/api/bot/checkin", {"token": member["token"], "checkpoint": "seminar",
                                       "by_telegram_id": "9001"})

        # Guruh mas'uli — faqat o'z guruhi.
        mine = self.client.post("/api/webapp/roster", json={
            "init_data": self.init_data(9001), "checkpoint": "seminar"}).get_json()
        self.assertEqual(mine["scope"]["group"], 1)
        self.assertEqual((mine["done"], mine["total"]), (1, 2))
        # Belgilanmaganlar tepada turadi.
        self.assertIsNone(mine["members"][0]["ts"])
        self.assertEqual(mine["members"][0]["id"], leader["id"])

        # Admin — hamma guruh, yig'indisi bilan.
        everyone = self.client.post("/api/webapp/roster", json={
            "init_data": self.init_data(555), "checkpoint": "seminar"}).get_json()
        self.assertEqual(everyone["total"], 3)
        self.assertEqual({g["group"]: (g["done"], g["total"]) for g in everyone["groups"]},
                         {1: (1, 2), 2: (0, 1)})

        # Oddiy a'zoga yopiq, imzosizga ham.
        self.assertEqual(self.client.post("/api/webapp/roster", json={
            "init_data": self.init_data(9100), "checkpoint": "seminar"}).status_code, 403)
        self.assertEqual(self.client.post("/api/webapp/roster", json={}).status_code, 401)

    def test_checkin_time_follows_the_configured_zone(self):
        admin = self.register("Panel Admin", "0000798")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        panel = self.panel(admin["id"])
        self.assertEqual(panel.post("/api/timezone", json={"offset": 5}).status_code, 200)
        self.assertEqual(panel.post("/api/timezone", json={"offset": 99}).status_code, 400)

        _, member, _ = self.build_group()
        stamped = self.scan(member["token"], "seminar").get_json()["ts"]
        expected = (datetime.datetime.now(datetime.timezone.utc)
                    + datetime.timedelta(hours=5)).strftime("%H:%M")
        self.assertEqual(stamped, expected)

    # ------------------------------------------------ Telegram guruh nazorati
    def seen_in_group(self, telegram_id, full_name, chat="-100777"):
        return self.post("/api/bot/group/seen", {
            "chat_id": chat, "telegram_id": telegram_id, "full_name": full_name,
            "status": "member", "source": "message"})

    def test_a_second_telegram_account_is_not_mistaken_for_a_stranger(self):
        """Odam bot bilan bir akkauntda tasdiqlab, guruhda boshqasida turishi mumkin."""
        leader, member, _ = self.build_group()
        self.sql("UPDATE participants SET fio='Abdulazizov Farrukh' WHERE id=?", member["id"])

        self.seen_in_group("9001", "Leader One")            # tasdiqlagan akkaunt
        self.seen_in_group("777001", "Farrux Abdulazizov")  # ikkinchi akkaunti
        self.seen_in_group("777002", "Acoustic Eshitish Markazi")   # begona

        audit = self.client.get("/api/bot/group/audit?chat_id=-100777",
                                headers=self.headers).get_json()
        self.assertEqual([x["id"] for x in audit["in_list"]], [leader["id"]])
        self.assertEqual([(x["id"], x["why"]) for x in audit["probable"]],
                         [(member["id"], "second_account")])
        self.assertEqual([x["telegram_id"] for x in audit["not_in_list"]], ["777002"])

    def test_an_unverified_participant_in_the_group_is_recognised_by_name(self):
        leader, member, _ = self.build_group()
        self.sql("UPDATE participants SET telegram_id=NULL WHERE id=?", member["id"])
        self.seen_in_group("777003", "Member One")
        audit = self.client.get("/api/bot/group/audit?chat_id=-100777",
                                headers=self.headers).get_json()
        self.assertEqual([(x["id"], x["why"]) for x in audit["probable"]],
                         [(member["id"], "not_verified")])
        self.assertEqual(audit["not_in_list"], [])

    def test_transliteration_variants_do_not_merge_different_people(self):
        for uz, other in (("Farrux", "Farrukh"), ("Toshxojaev", "Toshkhodjaev"),
                          ("Zhuraev", "Juraev"), ("Jumayev", "Jumaev")):
            self.assertTrue(server.word_matches(uz, other), f"{uz}={other}")
        for a, b in (("Karimov", "Karimova"), ("Shavkat", "Savkat"),
                     ("Abdullaev", "Abdullaeva")):
            self.assertFalse(server.word_matches(a, b), f"{a}!={b}")

    # ------------------------------------------------------- yosh va eksport
    def test_age_is_computed_and_visible_without_the_birth_date(self):
        leader, member, _ = self.build_group()
        born = datetime.date.today().replace(year=datetime.date.today().year - 30)
        # Tug'ilgan kunidan bir kun keyin — yosh to'liq 30.
        self.sql("UPDATE participants SET dob=? WHERE id=?",
                 (born - datetime.timedelta(days=1)).strftime("%d.%m.%Y"), member["id"])
        # Tug'ilgan kuni hali kelmagan — 29.
        self.sql("UPDATE participants SET dob=? WHERE id=?",
                 (born + datetime.timedelta(days=1)).strftime("%d.%m.%Y"), leader["id"])

        seen = {p["id"]: p for p in
                self.panel(leader["id"]).get("/api/bootstrap").get_json()["participants"]}
        self.assertEqual(seen[member["id"]]["age"], 30)
        self.assertEqual(seen[leader["id"]]["age"], 29)
        # Guruh mas'uli a'zoning yoshini ko'radi, lekin tug'ilgan sanasini emas.
        self.assertNotIn("dob", seen[member["id"]])
        self.assertIsNone(server.compute_age("xato"))
        self.assertIsNone(server.compute_age(""))

    def test_excel_export_follows_the_role_scope(self):
        leader, member, other = self.build_group()
        admin = self.register("Panel Admin", "0000795")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        import openpyxl, io as _io

        def sheet_of(response):
            self.assertEqual(response.status_code, 200)
            return openpyxl.load_workbook(_io.BytesIO(response.data)).active

        full = sheet_of(self.panel(admin["id"]).get("/api/export/participants.xlsx"))
        self.assertEqual(full.max_row - 1, 4)                 # hamma ishtirokchi
        self.assertIn("Pasport", [c.value for c in full[1]])

        # Guruh mas'uli o'z guruhini, pasport ustunlarisiz oladi.
        own = sheet_of(self.panel(leader["id"]).get("/api/export/participants.xlsx"))
        self.assertEqual(own.max_row - 1, 2)
        self.assertNotIn("Pasport", [c.value for c in own[1]])
        self.assertNotIn("Telefon", [c.value for c in own[1]])

        # A'zo faqat o'zini ko'radi; guruh eksporti unga umuman yopiq.
        mine = sheet_of(self.panel(member["id"]).get("/api/export/participants.xlsx"))
        self.assertEqual(mine.max_row - 1, 1)
        self.assertEqual(self.panel(member["id"]).get("/api/export/groups.xlsx").status_code, 403)
        self.assertEqual(self.client.get("/api/export/participants.xlsx").status_code, 401)

        # Guruhlar eksporti: har bir guruh alohida varaqda.
        book = openpyxl.load_workbook(_io.BytesIO(
            self.panel(admin["id"]).get("/api/export/groups.xlsx").data))
        self.assertEqual(book.sheetnames, ["1-Sazanchik", "2-Meduza"])

    # ------------------------------------------------------------- hujjatlar
    def upload(self, client, names, **extra):
        import io
        # Har fayl boshqacha mazmunda — aks holda takror deb hisoblanadi.
        data = {"files": [(io.BytesIO(b"%PDF-1.4 " + n.encode()), n) for n in names]}
        data.update(extra)
        return client.post("/api/docs/upload", data=data,
                           content_type="multipart/form-data")

    def test_uploaded_files_find_their_owner_by_name(self):
        leader, member, other = self.build_group()
        self.sql("UPDATE participants SET fio='Niyazov Bobir' WHERE id=?", other["id"])
        admin = self.register("Panel Admin", "0000782")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        panel = self.panel(admin["id"])

        body = self.upload(panel, [f"{member['id']} voucher.pdf",
                                   "Bilet_Niyazov_Bobir.pdf",
                                   "kimningdir_fayli.pdf"]).get_json()
        by_name = {d["file_name"]: d for d in body["saved"]}
        self.assertEqual(len(body["saved"]), 2)
        self.assertEqual(body["unmatched"], ["kimningdir_fayli.pdf"])
        self.assertEqual(by_name[f"{member['id']}_voucher.pdf"]["pid"], member["id"])
        self.assertEqual(by_name[f"{member['id']}_voucher.pdf"]["kind"], "voucher")
        self.assertEqual(by_name["Bilet_Niyazov_Bobir.pdf"]["pid"], other["id"])
        self.assertEqual(by_name["Bilet_Niyazov_Bobir.pdf"]["kind"], "ticket")

        # Fayl haqiqatan diskda va yuklab olinadi.
        doc_id = by_name[f"{member['id']}_voucher.pdf"]["id"]
        self.assertEqual(panel.get(f"/api/docs/file/{doc_id}").status_code, 200)

    def test_one_file_can_belong_to_several_people(self):
        """Uch kishilik xona voucherida uchalasining ismi bo'ladi."""
        leader, member, other = self.build_group()
        self.sql("UPDATE participants SET fio='Musaev Sardorjon' WHERE id=?", leader["id"])
        self.sql("UPDATE participants SET fio='Niyazov Bobir' WHERE id=?", member["id"])
        admin = self.register("Panel Admin", "0000786")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        server.DOCS_FILES = os.path.join(server.DOCS_DIR, "_files")
        panel = self.panel(admin["id"])

        body = self.upload(panel, ["Voucher_Musaev_Sardorjon_Niyazov_Bobir.pdf"]).get_json()
        self.assertEqual({s["pid"] for s in body["saved"]}, {leader["id"], member["id"]})
        # Disk ustida bitta nusxa, yozuvi ikkita.
        stored = {r[0] for r in self._all("SELECT stored FROM documents")}
        self.assertEqual(len(stored), 1)
        self.assertEqual(len(os.listdir(server.DOCS_FILES)), 1)

        # Ikkalasi ham o'z nusxasini oladi.
        for pid in (leader["id"], member["id"]):
            got = self.client.get(f"/api/bot/docs?id={pid}", headers=self.headers).get_json()
            self.assertEqual(len(got["documents"]), 1, pid)

        # Uchinchi odamga keyin qo'shish — faylni qayta yuklamasdan.
        doc_id = body["saved"][0]["id"]
        shared = panel.post("/api/docs/share", json={"id": doc_id, "pids": [other["id"]]})
        self.assertEqual(shared.get_json()["added"], [other["id"]])
        self.assertEqual(len(os.listdir(server.DOCS_FILES)), 1)

        # Bittasini o'chirsak fayl qolganlar uchun saqlanadi.
        panel.post("/api/docs/delete", json={"id": doc_id})
        self.assertEqual(len(os.listdir(server.DOCS_FILES)), 1)
        for row in self._all("SELECT id FROM documents"):
            panel.post("/api/docs/delete", json={"id": row[0]})
        self.assertEqual(os.listdir(server.DOCS_FILES), [])

    def make_pdf(self, pages):
        """Har sahifasida berilgan matn turgan PDF."""
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        import io
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=A4)
        for lines in pages:
            c.setFont("Helvetica", 12)
            y = 780
            for line in lines:
                c.drawString(60, y, line); y -= 20
            c.showPage()
        c.save()
        return buffer.getvalue()

    def upload_bytes(self, client, name, blob, **extra):
        import io
        data = {"files": (io.BytesIO(blob), name)}
        data.update(extra)
        return client.post("/api/docs/upload", data=data, content_type="multipart/form-data")

    def test_a_combined_pdf_is_split_so_each_person_gets_only_their_pages(self):
        leader, member, other = self.build_group()
        admin = self.register("Panel Admin", "0000788")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        server.DOCS_FILES = os.path.join(server.DOCS_DIR, "_files")
        panel = self.panel(admin["id"])

        blob = self.make_pdf([
            [f"Passenger: {leader['fio']}", "TAS - SSH"],
            ["Baggage: 20 kg"],                              # ismsiz davomi
            [f"Passenger: {member['fio']}", "TAS - SSH"],
            [f"Passenger: {other['fio']}", "TAS - SSH"],
        ])
        body = self.upload_bytes(panel, "hammasi.pdf", blob).get_json()
        self.assertEqual(body["split"], [{"file": "hammasi.pdf", "parts": 3}])
        self.assertEqual({s["pid"] for s in body["saved"]},
                         {leader["id"], member["id"], other["id"]})

        # Har kim faqat o'z bo'lagini oladi, va bo'laklar boshqa-boshqa fayl.
        stored = {r[0] for r in self._all("SELECT stored FROM documents")}
        self.assertEqual(len(stored), 3)

        from pypdf import PdfReader
        import io as _io
        pages_of = {}
        for row in self._all("SELECT pid,stored FROM documents"):
            with open(os.path.join(server.DOCS_FILES, row[1]), "rb") as fh:
                pages_of[row[0]] = PdfReader(_io.BytesIO(fh.read())).pages
        # Ismsiz sahifa oldingi odam bilan qoladi.
        self.assertEqual(len(pages_of[leader["id"]]), 2)
        self.assertEqual(len(pages_of[member["id"]]), 1)
        # Va birovning bo'lagida boshqasining ismi yo'q.
        self.assertNotIn(member["fio"], pages_of[leader["id"]][0].extract_text())

    def test_roommates_on_one_page_share_the_same_file_unsplit(self):
        leader, member, _ = self.build_group()
        admin = self.register("Panel Admin", "0000789")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        server.DOCS_FILES = os.path.join(server.DOCS_DIR, "_files")
        panel = self.panel(admin["id"])

        blob = self.make_pdf([[f"Guest 1: {leader['fio']}", f"Guest 2: {member['fio']}", "DBL"]])
        body = self.upload_bytes(panel, "room.pdf", blob).get_json()
        self.assertEqual(body["split"], [])                       # bo'linmaydi
        self.assertEqual({s["pid"] for s in body["saved"]}, {leader["id"], member["id"]})
        self.assertEqual(len(os.listdir(server.DOCS_FILES)), 1)   # bitta nusxa

    def test_a_pdf_without_a_text_layer_falls_back_to_the_filename(self):
        _, member, _ = self.build_group()
        admin = self.register("Panel Admin", "0000790")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        server.DOCS_FILES = os.path.join(server.DOCS_DIR, "_files")
        panel = self.panel(admin["id"])

        blob = self.make_pdf([["Hotel Rixos Sharm", "08-15 August 2026"]])   # ism yo'q
        body = self.upload_bytes(panel, f"{member['id']} voucher.pdf", blob).get_json()
        self.assertEqual([s["pid"] for s in body["saved"]], [member["id"]])
        self.assertEqual(body["split"], [])

    def test_owners_can_come_from_a_caption_when_the_filename_says_nothing(self):
        leader, member, _ = self.build_group()
        admin = self.register("Panel Admin", "0000787")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        server.DOCS_FILES = os.path.join(server.DOCS_DIR, "_files")
        panel = self.panel(admin["id"])

        body = self.upload(panel, ["IMG_2841.pdf"],
                           hint=f"{leader['id']}, {member['id']} chipta").get_json()
        self.assertEqual({s["pid"] for s in body["saved"]}, {leader["id"], member["id"]})
        self.assertEqual({s["kind"] for s in body["saved"]}, {"ticket"})

    def test_documents_are_only_visible_to_their_owner_and_the_staff(self):
        leader, member, other = self.build_group()
        admin = self.register("Panel Admin", "0000783")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        panel = self.panel(admin["id"])
        doc = self.upload(panel, [f"{member['id']}.pdf"]).get_json()["saved"][0]

        self.assertEqual(len(self.panel(member["id"]).get("/api/docs").get_json()["documents"]), 1)
        self.assertEqual(len(self.panel(leader["id"]).get("/api/docs").get_json()["documents"]), 1)
        # Boshqa guruh a'zosi na ro'yxatda ko'radi, na faylni ocha oladi.
        outsider = self.panel(other["id"])
        self.assertEqual(outsider.get("/api/docs").get_json()["documents"], [])
        self.assertEqual(outsider.get(f"/api/docs/file/{doc['id']}").status_code, 404)
        self.assertEqual(self.client.get(f"/api/docs/file/{doc['id']}").status_code, 401)

    def test_release_time_holds_documents_back_from_the_bot(self):
        _, member, _ = self.build_group()
        admin = self.register("Panel Admin", "0000784")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        panel = self.panel(admin["id"])
        self.upload(panel, [f"{member['id']} voucher.pdf", f"{member['id']} ticket.pdf"])

        ready = self.client.get(f"/api/bot/docs?id={member['id']}",
                                headers=self.headers).get_json()
        self.assertEqual(len(ready["documents"]), 2)   # vaqt qo'yilmagan — darrov tayyor

        later = (server._event_now().replace(tzinfo=None)
                 + datetime.timedelta(days=1)).strftime("%Y-%m-%dT%H:%M")
        self.assertEqual(panel.post("/api/docs/release",
                                    json={"release": {"ticket": later}}).status_code, 200)
        gated = self.client.get(f"/api/bot/docs?id={member['id']}",
                                headers=self.headers).get_json()
        self.assertEqual([d["kind"] for d in gated["documents"]], ["voucher"])
        self.assertEqual([d["kind"] for d in gated["pending"]], ["ticket"])

        # Ommaviy yuborish ham vaqti kelmaganini olmaydi.
        pending = self.client.get("/api/bot/docs/pending", headers=self.headers).get_json()
        self.assertEqual([d["kind"] for p in pending["people"] for d in p["documents"]],
                         ["voucher"])

    def test_name_spelling_variants_still_find_the_right_person(self):
        """Fayl nomini kim yozganiga qarab yozilishi har xil bo'ladi."""
        con = sqlite3.connect(server.DB_PATH)
        for pid, fio in (("ACO-901", "Jumaev Aziz"), ("ACO-902", "Karimov Kudratbek"),
                         ("ACO-903", "Karimova Nargiza"), ("ACO-904", "Mirzaev Shavkat")):
            con.execute("INSERT INTO participants(id,fio) VALUES(?,?)", (pid, fio))
        con.commit(); con.close()
        con = server.db()
        try:
            cases = {
                "JUMAYEV_AZIZBEK_TICKET.pdf": ["ACO-901"],   # yev→ev va +bek
                "JUMAEV_AZIZ_VOUCHER.pdf": ["ACO-901"],
                "MIRZAYEV_SHAVKAT.pdf": ["ACO-904"],
                "KARIMOV_KUDRATBEK.pdf": ["ACO-902"],
                # Karimov va Karimova — ikki xil odam, aralashib ketmasin.
                "KARIMOVA_NARGIZA.pdf": ["ACO-903"],
            }
            for name, expected in cases.items():
                self.assertEqual(server.match_participants(con, name), expected, name)
        finally:
            con.close()

    def test_organisers_are_left_out_of_rosters_and_statistics(self):
        leader, member, _ = self.build_group()
        staff = self.register("Organiser One", "0000793", "9500")
        self.sql("UPDATE participants SET staff=1 WHERE id=?", staff["id"])

        stats = self.client.get("/api/bot/stats", headers=self.headers).get_json()
        self.assertEqual(stats["total"], 3)              # 4 emas — tashkilotchi sanalmaydi
        everyone = self.client.get("/api/bot/recipients", headers=self.headers).get_json()
        self.assertNotIn(staff["id"], [r["id"] for r in everyone["recipients"]])

        # "Hammaga xabar" ham unga bormaydi.
        sent = self.post("/api/bot/message", {
            "from_telegram_id": "555", "scope": "all", "text": "salom"}).get_json()
        self.assertNotIn(staff["id"], [r["id"] for r in sent["recipients"]])

        # Uni check-in qilib bo'lmaydi — safarda qatnashmaydi.
        self.assertEqual(self.scan(staff["token"], "seminar").status_code, 403)

        # Hujjat hisobotida ham yo'q.
        state = self.client.get("/api/bot/docs/state", headers=self.headers).get_json()
        self.assertNotIn(staff["id"], [x["id"] for x in state["none"]])

    def test_a_note_travels_with_its_document_kind_in_the_right_language(self):
        """Reys vaqti o'zgargani kabi xabar chipta bilan birga ketadi."""
        _, member, _ = self.build_group()
        admin = self.register("Panel Admin", "0000794")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        server.DOCS_FILES = os.path.join(server.DOCS_DIR, "_files")
        panel = self.panel(admin["id"])
        self.upload(panel, [f"{member['id']} voucher.pdf", f"{member['id']} ticket.pdf"])

        self.assertEqual(panel.post("/api/docs/note", json={
            "kind": "ticket", "text": {"uz": "Reys vaqti o'zgardi",
                                       "ru": "Время рейса изменилось",
                                       "en": "The flight time has changed"}}).status_code, 200)

        got = self.client.get(f"/api/bot/docs?id={member['id']}",
                              headers=self.headers).get_json()
        self.assertEqual(got["notes"], {"ticket": "Reys vaqti o'zgardi"})   # tili yo'q -> uz

        self.sql("UPDATE participants SET lang='ru' WHERE id=?", member["id"])
        self.assertEqual(self.client.get(f"/api/bot/docs?id={member['id']}",
                                         headers=self.headers).get_json()["notes"],
                         {"ticket": "Время рейса изменилось"})

        # Voucherga izoh qo'yilmagan — u bilan hech nima ketmaydi.
        self.assertNotIn("voucher", got["notes"])

        # Chipta vaqti kelmagan bo'lsa izohi ham chiqmaydi.
        later = (server._event_now().replace(tzinfo=None)
                 + datetime.timedelta(days=1)).strftime("%Y-%m-%dT%H:%M")
        panel.post("/api/docs/release", json={"release": {"ticket": later}})
        held = self.client.get(f"/api/bot/docs?id={member['id']}",
                               headers=self.headers).get_json()
        self.assertEqual(held["notes"], {})

    def test_holding_delivery_stops_everything_going_out(self):
        """Fayllarni yuklab bo'lgunicha hech kimga yuborilmasin."""
        _, member, _ = self.build_group()
        admin = self.register("Panel Admin", "0000791")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        server.DOCS_FILES = os.path.join(server.DOCS_DIR, "_files")
        panel = self.panel(admin["id"])
        self.upload(panel, [f"{member['id']} voucher.pdf", f"{member['id']} ticket.pdf"])

        self.assertEqual(panel.post("/api/docs/release", json={"hold": True}).status_code, 200)
        held = self.client.get(f"/api/bot/docs?id={member['id']}",
                               headers=self.headers).get_json()
        self.assertEqual(held["documents"], [])
        self.assertEqual(len(held["pending"]), 2)
        self.assertEqual(self.client.get("/api/bot/docs/pending",
                                         headers=self.headers).get_json()["people"], [])
        self.assertTrue(panel.get("/api/docs").get_json()["hold"])

        # Ochilgach hammasi tayyor bo'ladi.
        panel.post("/api/docs/release", json={"hold": False})
        opened = self.client.get(f"/api/bot/docs?id={member['id']}",
                                 headers=self.headers).get_json()
        self.assertEqual(len(opened["documents"]), 2)

    def test_the_same_file_is_not_stored_twice_for_one_person(self):
        _, member, other = self.build_group()
        admin = self.register("Panel Admin", "0000792")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        server.DOCS_FILES = os.path.join(server.DOCS_DIR, "_files")
        panel = self.panel(admin["id"])

        name = f"{member['id']} voucher.pdf"
        first = self.upload(panel, [name]).get_json()
        self.assertEqual(len(first["saved"]), 1)

        again = self.upload(panel, [name]).get_json()
        self.assertEqual(again["saved"], [])
        self.assertEqual([d["pid"] for d in again["duplicates"]], [member["id"]])
        self.assertEqual(len(self._all("SELECT id FROM documents")), 1)
        self.assertEqual(len(os.listdir(server.DOCS_FILES)), 1)

        # Boshqa odamga o'sha fayl — yozuv qo'shiladi, fayl nusxalanmaydi.
        shared = self.upload(panel, [name], pid=other["id"]).get_json()
        self.assertEqual([s["pid"] for s in shared["saved"]], [other["id"]])
        self.assertEqual(len(os.listdir(server.DOCS_FILES)), 1)

    def test_sent_documents_are_not_sent_again(self):
        _, member, _ = self.build_group()
        admin = self.register("Panel Admin", "0000785")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        server.DOCS_DIR = os.path.join(self.tmp.name, "docs")
        doc = self.upload(self.panel(admin["id"]),
                          [f"{member['id']}.pdf"]).get_json()["saved"][0]

        pending = self.client.get("/api/bot/docs/pending", headers=self.headers).get_json()
        self.assertEqual(len(pending["people"]), 1)
        self.post("/api/bot/docs/sent", {"ids": [doc["id"]]})
        after = self.client.get("/api/bot/docs/pending", headers=self.headers).get_json()
        self.assertEqual(after["people"], [])
        # Odamning o'zi so'rasa baribir oladi — bir marta yuborilgani to'siq emas.
        self.assertEqual(len(self.client.get(
            f"/api/bot/docs?id={member['id']}", headers=self.headers).get_json()["documents"]), 1)

    def test_participant_page_and_scanner_stay_public(self):
        person = self.register("Public Person", "0000888")
        self.assertEqual(self.client.get(f"/api/p/{person['token']}").status_code, 200)
        self.assertEqual(self.client.get(f"/p/{person['token']}").status_code, 200)
        self.assertEqual(self.client.get("/scan").status_code, 200)

    def test_badge_template_keeps_multiple_logo_and_font_settings(self):
        admin = self.register("Badge Admin", "0009999")
        os.environ["PANEL_ADMIN_IDS"] = admin["id"]
        client = self.panel(admin["id"])
        badge = {"w": 70, "h": 110, "bg": "#fff", "elements": [
            {"id": "logo1", "type": "logo", "src": "data:image/png;base64,AAA", "x": 1, "y": 1, "w": 20},
            {"id": "logo2", "type": "logo", "src": "data:image/png;base64,BBB", "x": 30, "y": 1, "w": 20},
            {"id": "text1", "type": "text", "text": "Seminar", "font": "Georgia,serif",
             "italic": True, "lineHeight": 1.2, "letter": 0.4},
        ]}
        self.assertEqual(client.post("/api/badge", json=badge).status_code, 200)
        saved = client.get("/api/bootstrap").get_json()["badge"]
        self.assertEqual(saved["elements"][1]["src"], "data:image/png;base64,BBB")
        self.assertEqual(saved["elements"][2]["font"], "Georgia,serif")


if __name__ == "__main__":
    unittest.main()
